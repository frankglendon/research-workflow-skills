import copy
import json
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch
import openpyxl
from pptx import Presentation
from research_skills import workflows, demo
from research_skills.contracts import GateError, fingerprint


class WorkflowTests(unittest.TestCase):
    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory()
        self.addCleanup(self.temporary.cleanup)
        self.root = Path(self.temporary.name)
        self.data = demo.prepare(self.root / "inputs")

    def test_all_five_exports_are_valid_and_inputs_unchanged(self):
        before = [fingerprint(Path(p).read_bytes()) for p in self.data["inputs"]]
        results = demo.execute(self.data, self.root / "outputs")
        self.assertEqual(len(results), 5)
        self.assertTrue(all(r["openxml_errors"] == 0 for r in results))
        self.assertEqual(before, [fingerprint(Path(p).read_bytes()) for p in self.data["inputs"]])

    def test_qc_missing_review_blocks(self):
        plan = copy.deepcopy(self.data["qc_plan"])
        plan["reviews"].pop()
        demo.review(plan)
        with self.assertRaises(GateError):
            workflows.qc(self.data["survey"], plan, self.root / "blocked.xlsx")

    def test_post_review_edit_blocks(self):
        plan = copy.deepcopy(self.data["qc_plan"])
        plan["reviews"][0]["verdict"] = "issue"
        with self.assertRaisesRegex(GateError, "review"):
            workflows.qc(self.data["survey"], plan, self.root / "blocked.xlsx")

    def test_unconfirmed_binding_blocks(self):
        plan = copy.deepcopy(self.data["binding_plan"])
        plan["entries"][0]["confirmed"] = False
        demo.review(plan)
        with self.assertRaises(GateError):
            workflows.bind(self.data["data"], self.data["template"], plan, self.root / "blocked.pptx")

    def test_translation_number_change_blocks(self):
        plan = copy.deepcopy(self.data["translation_plan"])
        unit = next(key for key,value in plan["translations"].items() if "100" in value)
        plan["translations"][unit] = plan["translations"][unit].replace("100", "900")
        demo.review(plan)
        with self.assertRaises(GateError):
            workflows.translate(self.data["translation_source"], plan, self.root / "blocked.pptx")

    def test_questionnaire_jump_must_exist(self):
        plan = copy.deepcopy(self.data["questionnaire_plan"])
        plan["questions"][0]["jump_targets"] = ["missing"]
        demo.review(plan)
        with self.assertRaises(GateError):
            workflows.questionnaire(plan, self.root / "blocked.xlsx")

    def test_data_binding_writes_real_source_values(self):
        output = self.root / "bound.pptx"
        workflows.bind(self.data["data"], self.data["template"], self.data["binding_plan"], output)
        chart = next(s.chart for s in Presentation(output).slides[0].shapes if s.has_chart)
        self.assertEqual(list(chart.series[0].values), [100, 140])

    def test_qc_preserves_source_cells(self):
        output = self.root / "qc.xlsx"
        workflows.qc(self.data["survey"], self.data["qc_plan"], output)
        original = list(openpyxl.load_workbook(self.data["survey"]).active.values)
        self.assertEqual(original, list(openpyxl.load_workbook(output)["Survey"].values))

    def test_existing_output_is_never_overwritten(self):
        output = self.root / "existing.xlsx"
        output.write_bytes(b"keep this")
        with self.assertRaises(GateError):
            workflows.questionnaire(self.data["questionnaire_plan"], output)
        self.assertEqual(output.read_bytes(), b"keep this")

    def test_manifest_contains_no_source_answers_or_paths(self):
        output = self.root / "redacted.xlsx"
        result = workflows.qc(self.data["survey"], self.data["qc_plan"], output)
        receipt = json.loads(output.with_suffix(".xlsx.manifest.json").read_text())
        self.assertEqual(result, receipt)
        for forbidden in ["Useful service", "S001", str(self.root)]:
            self.assertNotIn(forbidden, json.dumps(receipt))

    def test_office_failure_blocks_artifact_and_manifest(self):
        output = self.root / "invalid.xlsx"
        with patch("research_skills.artifacts.validate_office", side_effect=GateError("Invalid Office file")):
            with self.assertRaises(GateError):
                workflows.questionnaire(self.data["questionnaire_plan"], output)
        self.assertFalse(output.exists())
        self.assertFalse(output.with_suffix(".xlsx.manifest.json").exists())
