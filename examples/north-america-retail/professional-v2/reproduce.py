"""Verify a frozen authored case without network or model calls.

This rechecks hashes, complete-answer routes, dictionary binding, workbook cells,
editable slide objects and Microsoft OOXML validity. It does not replay source
retrieval, semantic review, visual judgment, randomization or fieldwork.
"""
import argparse
import hashlib
import json
from pathlib import Path
import shutil
import tempfile
import zipfile
from xml.etree import ElementTree as ET

import openpyxl
from research_skills import survey_design as survey
from research_skills.artifacts import validate_office
from research_skills.contracts import GateError


def read(root, name):
    return json.loads((root/name).read_text(encoding="utf-8"))


def require(condition, message):
    if not condition:
        raise GateError(message)


def verify(root, office=True):
    root = Path(root).resolve()
    manifest = read(root, 'reviewed-inputs.json')
    for name, expected in manifest['sha256'].items():
        require(not Path(name).is_absolute() and '..' not in Path(name).parts,
                'Manifest paths must remain inside the case')
        target = (root/name).resolve()
        require(target.is_relative_to(root), 'Manifest symlink escapes the case')
        require(hashlib.sha256(target.read_bytes()).hexdigest() == expected,
                'Frozen case changed: '+name)
    spec = read(root, 'questionnaire-spec.json')
    codebook = read(root, 'codebook-plan.json')
    survey.require_review(spec)
    survey.require_review(codebook)
    survey.validate_codebook(spec, codebook)
    paths = survey.check_paths(spec)
    require(spec['routing_cases'] == read(root, 'routing-cases.json'), 'Routing files differ')
    data = read(root, 'workbook-data.json')
    for name, tabs in data.items():
        book = openpyxl.load_workbook(root/'deliverables'/(name+'.xlsx'), data_only=False)
        require(book.sheetnames == [t['name'] for t in tabs], 'Workbook tabs differ')
        for tab in tabs:
            sheet = book[tab['name']]
            rows = [tab['headers'], *tab['rows']]
            require(sheet.max_row == len(rows), 'Workbook row count differs')
            for i, row in enumerate(rows, 1):
                for j, value in enumerate(row, 1):
                    cell = sheet.cell(i, j)
                    expected = "'"+value if isinstance(value, str) and value.startswith('=') else value
                    require((cell.value if cell.value is not None else '') == expected,
                            f'Workbook cell differs: {name}/{tab["name"]}/{cell.coordinate}')
                    require(cell.data_type != 'f', 'Literal questionnaire content became a formula')
        book.close()
    sources = read(root, 'sources.json')
    claims = read(root, 'claims.json')
    pages = read(root, 'desk-research-content.json')
    source_ids = {s['id'] for s in sources}
    claim_by_id = {c['id']: c for c in claims}
    require(len(source_ids) == len(sources), 'Duplicate source IDs')
    require(len(claim_by_id) == len(claims), 'Duplicate claim IDs')
    for c in claims:
        require(set(c['source_ids']) <= source_ids, 'Unknown claim source')
        require(c['locator']['anchor_end'] > c['locator']['anchor_start'], 'Empty source locator')
        require(len(c['locator']['anchor_sha256']) == 64, 'Missing locator digest')
    for p in pages:
        require(set(p['source_ids']) <= source_ids, 'Unknown slide source')
        require(set(p['claim_ids']) <= set(claim_by_id), 'Unknown slide claim')
        for cid in p['claim_ids']:
            require(set(claim_by_id[cid]['source_ids']) <= set(p['source_ids']), 'Uncited slide claim')
    pptx = root/'deliverables/desk-research.pptx'
    ns = {'a':'http://schemas.openxmlformats.org/drawingml/2006/main',
          'c':'http://schemas.openxmlformats.org/drawingml/2006/chart'}
    with zipfile.ZipFile(pptx) as z:
        table_count = chart_count = 0
        for p in pages:
            xml = ET.fromstring(z.read(f'ppt/slides/slide{p["number"]}.xml'))
            text = ''.join(xml.itertext())
            require(p['title'] in text and p['decision'] in text, 'Slide narrative differs')
            tables = len(xml.findall('.//a:tbl', ns)); charts = len(xml.findall('.//c:chart', ns))
            require(bool(tables) == (p['kind'] == 'table'), 'Native table owner differs')
            require(bool(charts) == (p['kind'] == 'chart'), 'Native chart owner differs')
            table_count += tables; chart_count += charts
        embedded = [n for n in z.namelist() if n.startswith('ppt/embeddings/') and n.endswith('.xlsx')]
        require(len(embedded) == chart_count, 'Chart workbooks missing')
        if office:
            with tempfile.TemporaryDirectory() as tmp:
                for i, name in enumerate(embedded):
                    path = Path(tmp)/f'chart-{i}.xlsx'; path.write_bytes(z.read(name)); validate_office(path)
    if office:
        for name in ['questionnaire.xlsx', 'codebook.xlsx', 'desk-research.pptx']:
            validate_office(root/'deliverables'/name)
    return dict(questions=len(spec['questions']), respondent_questions=sum(q.get('role') != 'system' for q in spec['questions']),
                answer_variables=len(codebook['variables']), **paths,
                source_count=len(sources), claims=len(claims), slides=len(pages),
                native_tables=table_count, native_charts=chart_count,
                office_outputs=3, openxml_errors=0 if office else None,
                model_calls=0, network_calls=0, fieldwork_executed=False,
                semantic_review_replayed=False, visual_review_replayed=False)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--output', required=True)
    args = parser.parse_args()
    source = Path(__file__).resolve().parent
    result = verify(source)
    output = Path(args.output).resolve()
    require(not output.exists(), 'Choose a new replay output directory')
    output.mkdir(parents=True)
    for name in ['questionnaire.xlsx', 'codebook.xlsx', 'desk-research.pptx']:
        shutil.copy2(source/'deliverables'/name, output/name)
    (output/'validation.json').write_text(json.dumps(result, indent=2)+'\n', encoding='utf-8')
    print(json.dumps(result))


if __name__ == '__main__':
    main()
