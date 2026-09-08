"""Deterministic synthetic fixtures, not live model evaluations."""
import json
from pathlib import Path
import openpyxl
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches
from . import workflows
from .contracts import GateError, fingerprint
from .artifacts import validate_office
from .ooxml import normalize_xlsx, normalize_pptx
from .office import presentation, page, text, workbook_style


def review(plan):
    """Seal a pre-reviewed synthetic fixture; this function performs no review."""
    plan["reviewed"] = True
    plan["review_sha256"] = fingerprint({k:v for k,v in plan.items() if k != "review_sha256"})
    return plan


def prepare(folder):
    folder = Path(folder)
    folder.mkdir(parents=True, exist_ok=False)
    survey, data = folder / "survey.xlsx", folder / "metrics.xlsx"
    book = openpyxl.Workbook(); sheet = book.active; sheet.title = "Survey"
    for row in [("id", "score", "comment"), ("S001", 9, "Useful service"), ("S002", 14, ""), ("S003", 6, "Mixed experience")]:
        sheet.append(row)
    workbook_style(sheet); book.save(survey)
    book = openpyxl.Workbook(); sheet = book.active; sheet.title = "Metrics"
    for row in [("Period", "Revenue (synthetic units)"), ("Period A", 100), ("Period B", 140)]:
        sheet.append(row)
    workbook_style(sheet); book.save(data)
    template = folder / "chart-template.pptx"
    deck = presentation(); slide = page(deck, "Two reviewed cells update an editable chart")
    values = CategoryChartData(); values.categories = ["Period A", "Period B"]
    values.add_series("Revenue (synthetic units)", [0, 0])
    target = slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(0.9), Inches(2.2), Inches(11.4), Inches(4.3), values)
    deck.save(template)
    translation_source = folder / "slide-text.pptx"
    deck = presentation(); slide = page(deck, "Revenue was 100 units")
    text(slide, "A synthetic statement for translation validation.", 0.7, 2.6, 11.9, 1.3, 26)
    deck.save(translation_source)
    for path in [survey, data]: normalize_xlsx(path); validate_office(path)
    for path in [template, translation_source]: normalize_pptx(path); validate_office(path)
    digest = lambda paths: [fingerprint(Path(p).read_bytes()) for p in paths]
    qc_plan = review({"input_sha256": digest([survey]), "reviews": [
        {"row":2, "verdict":"clean", "note":"Synthetic review fixture"},
        {"row":3, "verdict":"issue", "note":"Invalid scale and missing comment"},
        {"row":4, "verdict":"uncertain", "note":"Keep uncertainty visible"}]})
    binding_plan = review({"input_sha256": digest([data,template]), "series_name":"Revenue (synthetic units)",
        "target":{"slide":0, "shape_id":target.shape_id}, "entries":[
            {"sheet":"Metrics", "cell":"B2", "label":"Period A", "confirmed":True},
            {"sheet":"Metrics", "cell":"B3", "label":"Period B", "confirmed":True}]})
    source_text = "ExampleCo reported 100 synthetic revenue units in Period A."
    evidence = [{"document_id":"synthetic-note", "quote":source_text}]
    ledger = {"documents":[{"id":"synthetic-note", "url":"https://example.org/synthetic-note",
        "text":source_text, "sha256":fingerprint(source_text.encode())}], "claims":[{
        "id":"claim-1", "statement":source_text, "status":"supported", "scope":{"kind":"synthetic"},
        "evidence":evidence, "evidence_sha256":fingerprint(evidence),
        "reviewed_statement_sha256":fingerprint(source_text.encode()), "semantic_reviewed":True}]}
    research_plan = review({"ledger":ledger})
    questionnaire_plan = review({"questions":[
        {"id":"Q1", "type":"single", "text":"Have you used this example service?", "options":[
            {"code":1, "label":"Yes"}, {"code":2, "label":"No"}], "jump_targets":["Q2", "END"]},
        {"id":"Q2", "type":"nps", "text":"How likely are you to recommend this example service?",
            "options":[{"code":i, "label":str(i)} for i in range(11)], "jump_targets":["END"]}]})
    spanish = {"Revenue was 100 units":"Los ingresos fueron de 100 unidades",
        "RESEARCH WORKFLOW SKILLS / SYNTHETIC DEMO":"RESEARCH WORKFLOW SKILLS / DEMOSTRACIÓN SINTÉTICA",
        "A synthetic statement for translation validation.":"Una afirmación sintética para validar la traducción.",
        "Synthetic data. No real organization or business result is represented.":"Datos sintéticos. No se representa ninguna organización ni resultado comercial real."}
    translations = {key:spanish[value] for key,value in workflows.translation_units(translation_source).items()}
    translation_plan = review({"input_sha256":digest([translation_source]), "translations":translations})
    result = dict(survey=survey, data=data, template=template, translation_source=translation_source,
        inputs=[survey,data,template,translation_source], qc_plan=qc_plan, binding_plan=binding_plan,
        research_plan=research_plan, questionnaire_plan=questionnaire_plan, translation_plan=translation_plan)
    for key,value in result.items():
        if key.endswith("_plan"):
            (folder/(key+".json")).write_text(json.dumps(value,indent=2)+"\n")
    return result


def execute(data, folder):
    folder = Path(folder); folder.mkdir(parents=True, exist_ok=True)
    return [workflows.qc(data["survey"], data["qc_plan"], folder/"01-survey-qc.xlsx"),
        workflows.research(data["research_plan"], folder/"02-evidence-report.pptx"),
        workflows.bind(data["data"], data["template"], data["binding_plan"], folder/"03-bound-chart.pptx"),
        workflows.questionnaire(data["questionnaire_plan"], folder/"04-questionnaire.xlsx"),
        workflows.translate(data["translation_source"], data["translation_plan"], folder/"05-translated-slide.pptx")]


def run(folder):
    folder = Path(folder)
    if folder.exists():
        raise GateError("Demo directory exists; choose a new output directory")
    data = prepare(folder/"inputs")
    results = execute(data, folder/"outputs")
    import copy
    bad = copy.deepcopy(data["research_plan"])
    bad["ledger"]["claims"][0]["evidence"][0]["quote"] = "Revenue was 900 units."
    review(bad)
    try:
        workflows.research(bad, folder/"outputs/blocked.pptx")
    except GateError:
        attack = "blocked"
    else:
        raise GateError("Fabricated quote was not blocked")
    summary = {"synthetic":True, "model_calls":0, "exports":len(results),
        "openxml_errors":sum(result["openxml_errors"] for result in results), "fabricated_quote":attack}
    (folder/"summary.json").write_text(json.dumps(summary,indent=2)+"\n")
    return summary
