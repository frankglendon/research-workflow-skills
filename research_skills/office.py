"""Plain native Office objects for synthetic demonstrations."""
from openpyxl.styles import Font, PatternFill
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor


def workbook_style(sheet):
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="17324D")
    for column in sheet.columns:
        sheet.column_dimensions[column[0].column_letter].width = min(65, max(18, max(len(str(c.value or "")) for c in column) + 3))


def presentation():
    deck = Presentation()
    deck.slide_width, deck.slide_height = Inches(13.333), Inches(7.5)
    deck.core_properties.author = "Research Workflow Skills"
    deck.core_properties.last_modified_by = "Research Workflow Skills"
    deck.core_properties.title = "Synthetic portfolio demonstration"
    return deck


def text(slide, value, left, top, width, height, size=24, color="17324D"):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    box.text_frame.word_wrap = True
    for index, line in enumerate(value.splitlines()):
        paragraph = box.text_frame.paragraphs[0] if index == 0 else box.text_frame.add_paragraph()
        paragraph.text = line
        paragraph.font.name = "Arial"
        paragraph.font.size = Pt(size)
        paragraph.font.color.rgb = RGBColor.from_string(color)
        paragraph.space_after = Pt(12)
    return box


def page(deck, title):
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    text(slide, "RESEARCH WORKFLOW SKILLS / SYNTHETIC DEMO", 0.7, 0.3, 12, 0.4, 12, "526779")
    text(slide, title, 0.7, 1.0, 12, 1.1, 32)
    text(slide, "Synthetic data. No real organization or business result is represented.", 0.7, 7.0, 12, 0.3, 11, "526779")
    return slide
