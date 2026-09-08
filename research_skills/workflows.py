"""Five intentionally narrow public workflows with executable contracts."""
from collections import Counter
from pathlib import Path
import math
import re
import zipfile
import openpyxl
from lxml import etree
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from .contracts import GateError, fingerprint, review_rows, validate_evidence
from .artifacts import export_artifact
from .office import workbook_style, presentation, page, text


def verify_plan(plan, sources=()):
    payload = {key:value for key,value in plan.items() if key != "review_sha256"}
    if plan.get("reviewed") is not True or plan.get("review_sha256") != fingerprint(payload):
        raise GateError("The current plan requires a completed review")
    if plan.get("input_sha256", []) != [fingerprint(Path(p).read_bytes()) for p in sources]:
        raise GateError("Input changed after review")


def qc(source, plan, output):
    verify_plan(plan, [source])
    book = openpyxl.load_workbook(source)
    sheet = book.active
    if list(next(sheet.values)) != ["id", "score", "comment"]:
        raise GateError("Public QC demo requires id, score, comment columns")
    reviews = review_rows(range(2, sheet.max_row + 1), plan.get("reviews", []))
    if "QC Review" in book.sheetnames:
        raise GateError("Input already contains a QC Review sheet")
    original = list(sheet.values)
    def build(path, run_id):
        result = book.create_sheet("QC Review")
        result.append(["Source row", "Rule findings", "Host verdict", "Review note"])
        for review in reviews:
            row = review["row"]
            score, comment = sheet.cell(row, 2).value, sheet.cell(row, 3).value
            findings = []
            if type(score) not in (int, float) or not math.isfinite(score) or not 0 <= score <= 10:
                findings.append("Score outside 0-10")
            if not str(comment or "").strip():
                findings.append("Empty comment")
            result.append([row, "; ".join(findings), review["verdict"], review.get("note", "")])
        workbook_style(result)
        book.save(path)
        if list(openpyxl.load_workbook(path)[sheet.title].values) != original:
            raise GateError("Source cells changed")
    return export_artifact(output, build, tool="survey-qc", inputs=[source], config=plan,
                           counts={"reviewed_rows": len(reviews)})


def research(plan, output):
    verify_plan(plan)
    counts = validate_evidence(plan["ledger"])
    ledger = plan["ledger"]
    documents = {d["id"]:d for d in ledger["documents"]}
    def build(path, run_id):
        deck = presentation()
        for claim in ledger["claims"]:
            slide = page(deck, "A claim with inspectable evidence")
            text(slide, claim["statement"], 0.7, 2.2, 11.9, 1.0, 28)
            quote = claim["evidence"][0]["quote"]
            text(slide, "SOURCE EXCERPT\n" + quote, 0.7, 3.7, 11.9, 1.4, 20)
            links = list(dict.fromkeys(documents[e["document_id"]]["url"] for e in claim["evidence"]))
            text(slide, "Sources: " + " | ".join(links), 0.7, 5.7, 11.9, 0.9, 14)
        deck.save(path)
    return export_artifact(output, build, tool="evidence-report", inputs=[], config=plan, counts=counts)


def bind(source, template, plan, output):
    verify_plan(plan, [source, template])
    entries = plan.get("entries", [])
    if not entries or any(e.get("confirmed") is not True for e in entries):
        raise GateError("Each source-to-chart mapping requires confirmation")
    if len({e["label"] for e in entries}) != len(entries):
        raise GateError("Chart categories must be unique")
    book = openpyxl.load_workbook(source, data_only=False)
    values = [book[e["sheet"]][e["cell"]].value for e in entries]
    if any(type(value) not in (int, float) or not math.isfinite(value) for value in values):
        raise GateError("Binding needs explicit finite numeric cells, not formulas or missing values")
    def build(path, run_id):
        deck = Presentation(template)
        target = plan["target"]
        matches = [s for s in deck.slides[target["slide"]].shapes if s.shape_id == target["shape_id"]]
        if len(matches) != 1 or not matches[0].has_chart:
            raise GateError("The reviewed chart target does not exist")
        if len(matches[0].chart.series) != 1:
            raise GateError("Public binding workflow supports one chart series")
        chart_data = CategoryChartData()
        chart_data.categories = [e["label"] for e in entries]
        chart_data.add_series(plan["series_name"], values)
        matches[0].chart.replace_data(chart_data)
        deck.save(path)
    return export_artifact(output, build, tool="data-binding", inputs=[source, template],
                           config=plan, counts={"mapped_cells": len(entries)})


def questionnaire(plan, output):
    verify_plan(plan)
    questions = plan.get("questions", [])
    ids = [q.get("id") for q in questions]
    if not ids or any(not value for value in ids) or len(set(ids)) != len(ids):
        raise GateError("Question IDs must be present and unique")
    for question in questions:
        if question.get("type") not in {"single", "open", "nps"} or not question.get("text"):
            raise GateError("Question requires text and a supported type")
        codes = [str(o.get("code", "")) for o in question.get("options", [])]
        if any(code in {"", "None"} for code in codes) or len(codes) != len(set(codes)):
            raise GateError("Option codes must be present and unique")
        if question["type"] == "single" and not codes:
            raise GateError("Single-choice questions need options")
        if question["type"] == "nps" and set(codes) != {str(i) for i in range(11)}:
            raise GateError("NPS requires the complete 0-10 scale")
        if any(t != "END" and t not in ids for t in question.get("jump_targets", [])):
            raise GateError("Jump target is missing")
    def build(path, run_id):
        book = openpyxl.Workbook()
        sheet = book.active
        sheet.title = "Questionnaire"
        sheet.append(["ID", "Type", "Question", "Options", "Jump targets"])
        for q in questions:
            options = "; ".join(f'{o["code"]}: {o["label"]}' for o in q.get("options", []))
            sheet.append([q["id"], q["type"], q["text"], options, ", ".join(q.get("jump_targets", []))])
        workbook_style(sheet)
        book.save(path)
    return export_artifact(output, build, tool="questionnaire", inputs=[], config=plan,
                           counts={"questions": len(questions)})


def translation_units(source):
    units = {}
    with zipfile.ZipFile(source) as archive:
        for name in sorted(archive.namelist()):
            if not re.fullmatch(r"ppt/slides/slide\d+\.xml", name):
                continue
            root = etree.fromstring(archive.read(name))
            for index,node in enumerate(root.findall(".//{http://schemas.openxmlformats.org/drawingml/2006/main}t")):
                units[f"{name}#{index}"] = node.text or ""
    return units


def translate(source, plan, output):
    verify_plan(plan, [source])
    units = translation_units(source)
    replacements = plan.get("translations", {})
    if not units or set(units) != set(replacements):
        raise GateError("Every slide text unit requires a reviewed translation")
    numbers = lambda value: Counter(re.findall(r"[-+]?\d+(?:[.,]\d+)*(?:%|％)?", value))
    for key,value in units.items():
        if not isinstance(replacements[key], str) or (value.strip() and not replacements[key].strip()):
            raise GateError("A translation is empty")
        if numbers(value) != numbers(replacements[key]):
            raise GateError("Translation changed protected numeric tokens")
    def build(path, run_id):
        with zipfile.ZipFile(source) as archive, zipfile.ZipFile(path, "w") as target:
            for item in archive.infolist():
                data = archive.read(item.filename)
                if re.fullmatch(r"ppt/slides/slide\d+\.xml", item.filename):
                    root = etree.fromstring(data)
                    for index,node in enumerate(root.findall(".//{http://schemas.openxmlformats.org/drawingml/2006/main}t")):
                        node.text = replacements[f"{item.filename}#{index}"]
                    data = etree.tostring(root, encoding="UTF-8", xml_declaration=True, standalone=True)
                target.writestr(item, data)
    return export_artifact(output, build, tool="slide-translation", inputs=[source], config=plan,
                           counts={"text_units": len(units)})
