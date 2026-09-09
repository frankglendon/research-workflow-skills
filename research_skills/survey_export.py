"""Office delivery for structured survey designs and codebooks."""
import json
from pathlib import Path
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from .artifacts import export_artifact
from .contracts import fingerprint
from . import survey_design as design


def _text(value):
    if isinstance(value, (list, dict)):
        return json.dumps(value, ensure_ascii=False) if value else ""
    return value


def _sheet(book, name, headers, rows):
    sheet = book.create_sheet(name)
    for row in [headers, *rows]:
        sheet.append([_text(v) for v in row])
    for row in sheet:
        for cell in row:
            if isinstance(cell.value, str):
                cell.data_type = "s"  # Survey content is text, never a spreadsheet formula.
            cell.font = Font(name="Arial", size=10, color="172B4D")
            cell.alignment = Alignment(vertical="top", wrap_text=True, horizontal="left", indent=1)
            cell.border = Border(bottom=Side(style="hair", color="D8E0EA"))
            if cell.row > 1 and cell.row % 2 == 0:
                cell.fill = PatternFill("solid", fgColor="F4F7FB")
            if cell.row == 1:
                cell.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="172B4D")
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A3
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.print_title_rows = "1:1"
    sheet.row_dimensions[1].height = 30
    for i, label in enumerate(headers, 1):
        width = 45 if any(t in label.lower() for t in ("text", "label", "rule", "definition", "decision", "missing", "logic")) else 23
        width = {"Code": 10, "Answer Code": 14, "Type": 14, "Population": 14, "Question ID": 16, "Module": 20, "Instruction": 35}.get(label, width)
        sheet.column_dimensions[openpyxl.utils.get_column_letter(i)].width = width
    for row in list(sheet)[1:]:
        sheet.row_dimensions[row[0].row].height = min(150, max(28, 15 * max((len(str(c.value or "")) // 28 + 1) for c in row)))
    return sheet


def _dictionary(book, variables):
    _sheet(book, "Variables", ["Variable ID", "Question ID", "Population", "Type", "Question label", "Base logic", "Missing rule", "Canonical ID"],
           [[v["variable"], v["question_id"], v["population"], v["type"], v["label"], v["base"], v["missing"], v["canonical_id"]] for v in variables])
    rows = []
    for v in variables:
        values = v["values"] or {"": "Text response"}
        if v["type"] == "numeric": values = {"": {"range": v["values"]["range"]}}
        for code, label in values.items():
            rows.append([v["question_id"], v["variable"], v["type"], code, v["label"], label, v["population"]])
    _sheet(book, "Datamap", ["Question ID", "Variable ID", "Type", "Answer Code", "Question Label", "Answer Label", "Population"], rows)


def questionnaire(spec, output):
    counts = design.validate(spec); design.require_review(spec)
    counts.update(design.check_paths(spec))
    variables = design.variables(spec)
    def build(path, run_id):
        book = openpyxl.Workbook(); book.remove(book.active)
        _sheet(book, "Study", ["Field", "Value"], [[k, v] for k, v in spec["meta"].items()] + [["QC plan", spec["qc_plan"]], ["Specification SHA256", design.spec_hash(spec)]])
        _sheet(book, "Sampling", ["Population", "Label", "Sampling specification"], [[p["key"], p.get("name", ""), p.get("sampling", {})] for p in spec["populations"]])
        _sheet(book, "Analysis plan", ["Objective", "Decision", "Metric", "Questions"],
               [[o["id"], o["decision"], o["metric"], [q["population"] + ":" + q["id"] for q in spec["questions"] if o["id"] in q["objective_ids"]]] for o in spec["objectives"]])
        rows, programming = [], []
        for q in spec["questions"]:
            base = [q["population"], q["module"], q["id"], q["type"]]
            rows.append(base + [q["text"], "", "", q.get("instruction", "")])
            rows.extend(["", "", "", "", "", r["code"], r["label"], "Matrix row"] for r in q.get("rows", []))
            for option in q.get("options", []):
                notes = [title for key, title in [("specify", "Collect other text"), ("exclusive", "Exclusive"), ("fixed", "Fixed position")] if option.get(key)]
                if option.get('show_if'): notes.append('Display: ' + _text(option['show_if']))
                rows.append(["", "", "", "", "", option["code"], option["label"], "; ".join(notes)])
            programming.append([q["population"], q["id"], q.get("show_if", []), q.get("routes", []), q.get("default", "NEXT"), q.get("randomize", False),
                                {k: q[k] for k in ("min_selections", "max_selections", "rank_count", "min", "max", "options_from", "rows_from") if k in q}, q.get("programming_note", "")])
        _sheet(book, "Questionnaire", ["Population", "Module", "Question ID", "Type", "Question text", "Code", "Answer label", "Instruction"], rows)
        _sheet(book, "Programming", ["Population", "Question ID", "Display logic", "Route logic", "Default", "Randomize", "Limits", "Programming note"], programming)
        _dictionary(book, variables)
        book.save(path)
    return export_artifact(output, build, tool="survey-design", inputs=[], config=spec, counts={**counts, "variables": len(variables)})


def codebook(spec, plan, output):
    counts = design.validate_codebook(spec, plan)
    design.require_review(spec); design.require_review(plan)
    design.check_paths(spec)
    def build(path, run_id):
        book = openpyxl.Workbook(); book.remove(book.active)
        _sheet(book, "Version", ["Field", "Value"], [["Codebook version", plan["version"]], ["Questionnaire version", spec["meta"]["version"]], ["Questionnaire SHA256", design.spec_hash(spec)]])
        _dictionary(book, plan["variables"])
        rows = []
        for frame in plan["codeframes"]:
            for code in frame["codes"]:
                rows.append([frame["population"], frame["question_id"], frame["unit"], frame["mode"], code["code"], code.get("parent", ""), code["label"], code["definition"], code["include"], code["exclude"], code.get("special", False), code.get("exclusive", False)])
        _sheet(book, "Open coding", ["Population", "Question ID", "Unit", "Mode", "Code", "Parent", "Label", "Definition", "Include rule", "Exclude rule", "Special", "Exclusive"], rows)
        book.save(path)
    return export_artifact(output, build, tool="codebook", inputs=[], config={"questionnaire": spec, "codebook": plan}, counts=counts)


def inspect_datamap(source):
    """Read three common dictionary layouts. Return review candidates, never infer unknown layouts."""
    book = openpyxl.load_workbook(source, read_only=True, data_only=False)
    layouts = []
    try:
        for sheet in book:
            rows = sheet.iter_rows(values_only=True)
            for n in range(1, 11):
                row = next(rows, None)
                if row is None: break
                labels = [str(v or "").strip() for v in row]
                if "Question ID" in labels and "Variable ID" in labels and "Answer Code" in labels:
                    kind = "long_datamap"
                elif "Fields" in labels and "Label" in labels and "Type" in labels:
                    kind = "field_metadata"
                elif labels[:3] == ["值", "", "标签"]:
                    kind = "value_labels"
                else: continue
                layouts.append({"sheet": sheet.title, "header_row": n, "layout": kind,
                                "columns": [{"index": i + 1, "label": v} for i, v in enumerate(labels)],
                                "data_rows": max(0, sheet.max_row - n)})
                break
    finally:
        book.close()
    design.need(layouts, "No supported Datamap header found; inspect the workbook and map columns explicitly")
    return {"source_sha256": fingerprint(Path(source).read_bytes()), "layouts": layouts,
            "status": "needs_mapping_review", "note": "Structural inventory only; formulas, multi-response storage and loops require review before import."}
